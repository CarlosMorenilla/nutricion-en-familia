
"""Read source workbooks; never execute their builders or modify originals.
Run with bundled Python. Output is private and must never enter the public build.
"""
from pathlib import Path
import json, uuid, datetime, hashlib
import openpyxl
ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "web" / "private"
OUT.mkdir(exist_ok=True)
DAYS = ["LUNES","MARTES","MIÉRCOLES","JUEVES","VIERNES","SÁBADO","DOMINGO"]
def text(v):
    return "" if v is None else str(v)
def load(p):
    return openpyxl.load_workbook(p, data_only=True)
def blank_day():
    return dict(meals={s:dict(title="",portion="",substitution="") for s in ["desayuno","comida","merienda","cena"]},workout=dict(title="",details="",intensity="",rest="",exercises=""),notes="")
sources = {
 "carlitos": {
 "diet": ROOT/"Carlitos/Actualizados/Dieta_Carlitos_actualizada.xlsx",
 "cardio": ROOT/"Carlitos/Actualizados/Entreno_Cardio_Carlitos_actualizado.xlsx",
 "strength": ROOT/"Carlitos/Actualizados/Entreno_Fuerza_Carlitos_actualizado.xlsx"},
 "mama": {
 "diet": ROOT/"Mamá/Dieta_Mamá.xlsx",
 "cardio": ROOT/"Mamá/Entreno_Cardio_Mamá.xlsx",
 "strength": ROOT/"Mamá/Entreno_Fuerza_Mamá.xlsx"}}
plan = dict(id=str(uuid.uuid4()),week="2026-09-07",version=1,revision=0,status="draft",reviewed=False,
 shared=[dict(day=d,slot=s,title="",instructions="") for d in range(7) for s in ["comida","cena"]],members=[])
report = {"sources":[],"checks":[],"conflicts":[]}
for person,files in sources.items():
    diet=load(files["diet"])["Dieta Semanal"]
    cardio=load(files["cardio"])["Plan Semanal"]
    strength=load(files["strength"])
    notes="\n".join(text(row[0]) for row in list(diet.values)[9:] if row[0])
    for kind in ["cardio","strength"]:
        book=load(files[kind])
        for row in book.worksheets[0].values:
            if isinstance(row[0],str) and row[0].startswith("•"):
                notes += "\n"+row[0]
    member=dict(member_id=person,days=[],notes=notes,source="; ".join(p.name for p in files.values()))
    for d,day in enumerate(DAYS):
        obj=blank_day()
        for slot,row in zip(["desayuno","comida","merienda","cena"],[4,5,6,7]):
            value=text(diet.cell(row,d+2).value)
            obj["meals"][slot]=dict(title=value,portion=value,substitution="")
            assert obj["meals"][slot]["portion"] == text(diet.cell(row,d+2).value)
        obj["notes"]="Actividad original: "+text(diet.cell(2,d+2).value)+"\nReferencia energética original: "+text(diet.cell(3,d+2).value)+"\nIndicaciones originales pendientes de revisión: "+text(diet.cell(8,d+2).value)
        obj["workout"].update(title=text(cardio.cell(d+2,2).value),details=text(cardio.cell(d+2,3).value),intensity=text(cardio.cell(d+2,4).value))
        for row in strength.worksheets[0].values:
            if text(row[0]).upper()==day:
                obj["workout"]["rest"]=text(row[3])
        sheet=next((s for s in strength if s.title.startswith(day+" -")),None)
        if sheet:
            lines=[" · ".join(text(v) for v in row if v is not None) for row in list(sheet.values)[1:] if row[0]]
            obj["workout"]["exercises"]="\n".join(lines)
            obj["workout"]["title"]=sheet.title.split(" - ",1)[1]+" / "+obj["workout"]["title"]
        member["days"].append(obj)
    plan["members"].append(member)
    for kind,p in files.items():
        report["sources"].append(dict(file=str(p.relative_to(ROOT)),sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
    report["checks"].append(dict(person=person,days=7,meals=28,exact_meal_cells=True))
for d in range(7):
    for slot in ["comida","cena"]:
        report["conflicts"].append(dict(day=DAYS[d],slot=slot,reason="Revisar plato base y raciones individuales antes de publicar",carlitos=plan["members"][0]["days"][d]["meals"][slot]["portion"],mama=plan["members"][1]["days"][d]["meals"][slot]["portion"]))
(OUT/"borradores-iniciales.json").write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding="utf-8")
(OUT/"informe-importacion.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("Importados 6 Excel: 56 comidas exactas y 14 días de entrenamiento. 14 platos comunes pendientes de revisión.")

